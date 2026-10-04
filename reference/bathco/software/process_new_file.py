# Reads a new Excel file dropped into Desktop/bathco-data
# and appends a structured digest to the Claude memory file.
import sys
import os
import openpyxl
from datetime import datetime

MEMORY_FILE = r"C:\Users\1st Choice\.claude\projects\C--Users-1st-Choice\memory\bathco-new-data.md"
MAX_ROWS = 300

def read_sheet(ws):
    rows = []
    for row in ws.iter_rows(values_only=True):
        if all(v is None for v in row):
            continue
        rows.append("\t".join(str(v) if v is not None else "" for v in row))
        if len(rows) >= MAX_ROWS:
            rows.append(f"[... truncated at {MAX_ROWS} rows ...]")
            break
    return "\n".join(rows)

def process_file(path):
    filename = os.path.basename(path)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        sheets_content = []
        for name in wb.sheetnames:
            ws = wb[name]
            content = read_sheet(ws)
            if content.strip():
                sheets_content.append(f"### Sheet: {name}\n{content}")
        full_content = "\n\n".join(sheets_content)
    except Exception as e:
        full_content = f"ERROR reading file: {e}"

    entry = f"""
---
## New File Detected: {filename}
**Added:** {timestamp}
**Path:** {path}

{full_content}

---
"""
    # Ensure memory file exists with frontmatter
    if not os.path.exists(MEMORY_FILE):
        header = """---
name: bathco-new-data
description: New Excel files dropped into C:\\Users\\1st Choice\\Desktop\\bathco-data\\ — auto-processed and appended here for Claude to read in future conversations
metadata:
  type: project
---

# Bathco New Data Inbox

Files are automatically read when dropped into the watched folder.

"""
        with open(MEMORY_FILE, "w", encoding="utf-8") as f:
            f.write(header)

    with open(MEMORY_FILE, "a", encoding="utf-8") as f:
        f.write(entry)

    print(f"[OK] Processed: {filename} -> appended to memory at {timestamp}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python process_new_file.py <path_to_xlsx>")
        sys.exit(1)
    process_file(sys.argv[1])
