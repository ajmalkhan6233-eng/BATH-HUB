import openpyxl
import os
import sys

sys.stdout = open("C:/Users/1st Choice/bathco-data/excel_dump.txt", "w", encoding="utf-8")

def read_excel(path, max_rows=200):
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        out = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            out.append(f"\n=== SHEET: {sheet_name} ===")
            rows_read = 0
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                out.append("\t".join(str(v) if v is not None else "" for v in row))
                rows_read += 1
                if rows_read >= max_rows:
                    out.append(f"[... truncated at {max_rows} rows ...]")
                    break
        return "\n".join(out)
    except Exception as e:
        return f"ERROR reading {path}: {e}"

files = [
    r"C:\Users\1st Choice\bathco-data\BATHCO COMPLETE BUSINESS REPORT 29MAY2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\GRN.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\NON MOVING.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\PRICE LIST.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\QUANTITY AND PRICE.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\SLOW.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\STOCK FAST.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\T SALE.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\4.22\4.22\TOTAL Q.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\DALY SALES\DALY SALES\DALY SALES REPORT.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\DALY SALES\DALY SALES\01.02.26.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\DALY SALES\DALY SALES\05.02.26.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\New folder (3)\New folder (3)\DALY SALES REPORT.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\New folder (3)\New folder (3)\26.01.31.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\TOTAL.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\BATHCO_SIMPLE_REPORT.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\AFDAS.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\BFNH.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\CFBDF.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\FNGJN.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\OUTBOUND_CHEQ_CLEAN.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\SFSGS.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\VGNGF.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\12-05\12-05-2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\13-05\13-05-2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\14-05\14-05-2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\15-05\15-05-2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\16-05\16-05-2026.xlsx",
    r"C:\Users\1st Choice\bathco-data\extracted\URGENT\URGENT\19-05-2026\19-05-2026.xlsx",
]

for f in files:
    if os.path.exists(f):
        print(f"\n{'='*60}")
        print(f"FILE: {os.path.basename(f)}")
        print(f"PATH: {f}")
        print('='*60)
        print(read_excel(f))
    else:
        print(f"\nMISSING: {f}")
